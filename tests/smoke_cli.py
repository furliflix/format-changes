# -*- coding: utf-8 -*-
"""自检脚本：生成测试文件，把各个转换引擎真实跑一遍。

用法：
    python tests/smoke_cli.py

打包脚本和 GitHub Actions 都会先跑它，确认转换链路没问题。
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.core.deps import run_capture  # noqa: E402
from app.core.engine import Converter  # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, fn) -> None:
    try:
        fn()
        RESULTS.append((name, True, ""))
        print(f"[✓] {name}")
    except Exception as e:
        RESULTS.append((name, False, str(e)))
        print(f"[✗] {name}：{e}")


def main() -> int:
    conv = Converter()
    deps = conv.deps
    print("运行环境：")
    for name, ok, desc in deps.summary_lines():
        print(f"    [{'OK' if ok else '--'}] {name}：{desc}")
    print()

    tmp = Path(tempfile.mkdtemp(prefix="fw-smoke-"))
    try:
        fixture_audio = tmp / "声音测试.wav"
        fixture_video = tmp / "视频测试.mkv"
        fixture_image = tmp / "图片测试.png"
        fixture_docx = tmp / "文档测试.docx"
        fixture_pdf = tmp / "PDF测试.pdf"

        # ---------------- 生成测试素材 ----------------
        if deps.ffmpeg.ok:
            def make_audio():
                r = run_capture([
                    deps.ffmpeg.path, "-hide_banner", "-y",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=1.5",
                    str(fixture_audio),
                ], timeout=60)
                assert fixture_audio.exists() and fixture_audio.stat().st_size > 0, r.stdout
            check("生成测试音频", make_audio)

            def make_video():
                r = run_capture([
                    deps.ffmpeg.path, "-hide_banner", "-y",
                    "-f", "lavfi", "-i", "testsrc2=duration=2:size=320x240:rate=15",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                    "-c:v", "libx264" if (not deps.ffmpeg.encoders or "libx264" in deps.ffmpeg.encoders) else "mpeg4",
                    "-c:a", "aac", "-shortest", str(fixture_video),
                ], timeout=120)
                assert fixture_video.exists() and fixture_video.stat().st_size > 0, r.stdout
            check("生成测试视频", make_video)

        def make_image():
            from PIL import Image
            Image.new("RGB", (400, 260), (46, 107, 230)).save(fixture_image)
        if deps.pillow:
            check("生成测试图片", make_image)

        def make_docx():
            from docx import Document
            doc = Document()
            doc.add_heading("自检测试文档", level=1)
            doc.add_paragraph("这是一段中文测试文字 with English words。")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "左上"
            table.cell(0, 1).text = "右上"
            table.cell(1, 0).text = "左下"
            table.cell(1, 1).text = "右下"
            doc.save(str(fixture_docx))
        if deps.python_docx:
            check("生成测试 Word 文档", make_docx)

        # ---------------- 转换测试 ----------------
        def t(src: Path, target: str, quality: str = "fast") -> None:
            assert src.exists(), f"测试素材不存在：{src}"
            out = conv.convert(src, target, dst_dir=tmp / "out", quality=quality)
            assert out.exists() and out.stat().st_size > 0, f"输出文件为空：{out}"

        if deps.ffmpeg.ok and fixture_audio.exists():
            check("wav → mp3", lambda: t(fixture_audio, "mp3"))
            check("wav → ogg", lambda: t(fixture_audio, "ogg"))
            check("wav → flac", lambda: t(fixture_audio, "flac"))

        if deps.ffmpeg.ok and fixture_video.exists():
            check("mkv → mp4", lambda: t(fixture_video, "mp4"))
            check("mkv → webm", lambda: t(fixture_video, "webm", "small"))
            check("mkv → 提取 mp3", lambda: t(fixture_video, "mp3"))
            check("mkv → gif", lambda: t(fixture_video, "gif", "small"))

        if deps.pillow and fixture_image.exists():
            check("png → jpg", lambda: t(fixture_image, "jpg"))
            check("png → webp", lambda: t(fixture_image, "webp"))
            check("png → pdf", lambda: t(fixture_image, "pdf"))

        if fixture_docx.exists():
            def docx_pdf():
                out = conv.convert(fixture_docx, "pdf", dst_dir=tmp / "out", quality="fast")
                assert out.exists() and out.stat().st_size > 0
                shutil.copy(out, fixture_pdf)
            check("docx → pdf（LibreOffice 或基础模式）", docx_pdf)
            check("docx → txt", lambda: t(fixture_docx, "txt"))

        if deps.pymupdf and fixture_pdf.exists():
            check("pdf → txt", lambda: t(fixture_pdf, "txt"))

        # ---------------- 汇总 ----------------
        print()
        failed = [r for r in RESULTS if not r[1]]
        print(f"自检完成：{len(RESULTS) - len(failed)} 通过，{len(failed)} 失败")
        if failed:
            for name, _, err in failed:
                print(f"  失败：{name} -> {err.splitlines()[0] if err else ''}")
            return 1
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
