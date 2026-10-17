# -*- coding: utf-8 -*-
"""PDF 后端（PyMuPDF）：PDF → 文本 / Word。"""
from __future__ import annotations

from pathlib import Path

from ..cancel import CancelToken
from ..deps import Deps
from ..errors import CancelledError, ConvertError, MissingDependency


class PdfBackend:
    def __init__(self, deps: Deps):
        self.deps = deps

    def convert(
        self,
        src: Path,
        target: str,
        tmp_out: Path,
        q,
        progress,
        cancel: CancelToken,
    ) -> None:
        if not self.deps.pymupdf:
            raise MissingDependency("处理 PDF 需要 PyMuPDF：pip install PyMuPDF")
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:  # 兼容老版本 PyMuPDF
            import fitz  # type: ignore

        try:
            doc = fitz.open(str(src))
        except Exception as e:
            raise ConvertError(f"无法打开 PDF：{e}") from e
        try:
            n = doc.page_count
            if n == 0:
                raise ConvertError("这个 PDF 没有任何页面。")
            if target == "txt":
                parts = []
                for i in range(n):
                    if cancel.cancelled():
                        raise CancelledError()
                    parts.append(doc.load_page(i).get_text("text"))
                    progress(0.02 + 0.95 * (i + 1) / n, f"提取文本 {i + 1}/{n} 页")
                text = "\n\n".join(p.rstrip() for p in parts).strip() + "\n"
                tmp_out.write_text(text, encoding="utf-8")
            elif target == "docx":
                if not self.deps.python_docx:
                    raise MissingDependency("PDF 转 Word 需要 python-docx：pip install python-docx")
                from docx import Document
                from docx.shared import Pt, RGBColor

                document = Document()
                note = document.add_paragraph()
                run = note.add_run("（由 Format Changes 基础模式转换，复杂排版可能丢失）")
                run.italic = True
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

                for i in range(n):
                    if cancel.cancelled():
                        raise CancelledError()
                    page = doc.load_page(i)
                    blocks = [b for b in page.get_text("blocks") if (b[4] or "").strip()]
                    blocks.sort(key=lambda b: (round(b[1] / 6.0), b[0]))
                    for b in blocks:
                        document.add_paragraph(str(b[4]).strip())
                    if i != n - 1:
                        document.add_page_break()
                    progress(0.02 + 0.95 * (i + 1) / n, f"整理内容 {i + 1}/{n} 页")
                document.save(str(tmp_out))
            else:
                raise ConvertError(f"不支持的 PDF 目标格式：{target}")
        finally:
            doc.close()
        progress(1.0, "完成")
